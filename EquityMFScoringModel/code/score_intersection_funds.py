"""
Builds and scores the Equity mutual fund universe restricted to the 9
scoring parameters that actually have real data in scheme_masters_1.json
and risk.json, keeping only funds where EVERY one of those 9 parameters
is present (the intersection):

    Rolling Returns (3Y/5Y)       (blended 3Y/5Y return, scored as excess vs.
                                    the fund's category peer average -- there is
                                    no benchmark-index return series in the data,
                                    so category average stands in for "benchmark")
    Alpha (3Y)
    CAGR vs Category Avg (5Y)     (5Y return only, scored as excess vs. the same
                                    category peer average, on its own 5Y-only
                                    baseline -- related to Rolling Returns but not
                                    the same calculation; no 10Y data exists anywhere)
    Standard Deviation (vs Category)
    Sharpe Ratio (3Y)
    Sortino Ratio (3Y)
    Beta
    Exit Load Structure
    AUM Size (Category-adjusted)

Everything else from Tanisha Mam's scoring matrix (CAGR 10Y, Fund Manager
Tenure, Investment Style Consistency, Portfolio Concentration, Sector
Diversification, Expense Ratio, ...) has zero data source across the
project's JSON files for Equity funds and is left out entirely -- not
scored as missing, just not part of this matrix.

Every parameter is scored using the EXACT 1-5 bands defined in
"equity_mf_scoring_matrix by Tanisha Mam (1).xlsx" (its hidden "Sheet1"
carries the full 1/2/3/4/5 breakpoints for Alpha, CAGR, Standard
Deviation, Sharpe, Sortino and Beta; "Rating Scale & Legend" carries the
category-specific AUM Cr ranges and the final composite rating table) --
not a generic z-score or percentile calculation. Where the source file
doesn't have a real benchmark-return series to compare against, the
fund's own category peer average is used as the benchmark/category-average
proxy her bands are defined against (documented inline at each spot).

Step 1 (build): filters scheme_masters_1.json + risk.json down to
Equity-category funds that have all 9 fields, and writes that combined
intersection dataset to equity_fund_intersection_dataset.json/.xlsx.

Step 2 (score): scores every fund in that intersection dataset against
the 9 parameters (weights renormalized from Tanisha Mam's matrix to sum
to 1.0 over just these 9) and writes the ranked result to
equity_fund_intersection_scores.json/.xlsx.

Run: python score_intersection_funds.py [--top N]
"""

import argparse
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).parent.parent
DATA_DIR = BASE_DIR / "data"
OUTPUT_DIR = BASE_DIR / "output"

# weights taken from the master scoring matrix for just these 9 parameters,
# renormalized (raw weights summed to 0.60) so they sum to 1.0
WEIGHTS = {
    "rolling_returns_vs_benchmark": 0.10 / 0.60,
    "alpha_3y": 0.08 / 0.60,
    "cagr_vs_category_avg": 0.07 / 0.60,
    "sd_vs_category": 0.08 / 0.60,
    "sharpe_3y": 0.08 / 0.60,
    "sortino_3y": 0.07 / 0.60,
    "beta": 0.04 / 0.60,
    "exit_load_structure": 0.03 / 0.60,
    "aum_size_category_adjusted": 0.05 / 0.60,
}

PARAM_LABELS = {
    "rolling_returns_vs_benchmark": "Rolling Returns (3Y/5Y)",
    "alpha_3y": "Alpha (3Y)",
    "cagr_vs_category_avg": "CAGR vs Category Avg (5Y)",
    "sd_vs_category": "Standard Deviation (vs Category)",
    "sharpe_3y": "Sharpe Ratio (3Y)",
    "sortino_3y": "Sortino Ratio (3Y)",
    "beta": "Beta",
    "exit_load_structure": "Exit Load Structure",
    "aum_size_category_adjusted": "AUM Size (Category-adjusted)",
}

PARAM_CATEGORY = {
    "rolling_returns_vs_benchmark": "Return",
    "alpha_3y": "Return",
    "cagr_vs_category_avg": "Return",
    "sd_vs_category": "Risk",
    "sharpe_3y": "Risk",
    "sortino_3y": "Risk",
    "beta": "Risk",
    "exit_load_structure": "Cost & Operational",
    "aum_size_category_adjusted": "Cost & Operational",
}

# composite-score interpretation, taken verbatim from the matrix's
# "Rating Scale & Legend" sheet -> "TOTAL WEIGHTED SCORE -- INTERPRETATION"
RATING_BANDS = [
    (4.50, "Outstanding", "Strong Buy"),
    (4.00, "Excellent", "Buy"),
    (3.50, "Good", "Buy / Watch"),
    (3.00, "Average", "Hold / Consider"),
    (2.50, "Below Average", "Avoid"),
    (0.00, "Poor", "Reject"),
]

# AUM Size Cr ranges, per fund category, from the matrix's "Rating Scale &
# Legend" sheet -> "AUM SIZE SCORING GUIDE BY FUND CATEGORY". score1_max is
# the "too small" ceiling; score3/score5 ranges deliberately overlap in the
# source file itself (e.g. Mid-Cap's score-3 range runs 1,000-10,000 Cr while
# its score-5 "ideal" range is 2,000-15,000 Cr) -- treated below as: inside
# the score-5 range wins first, then score-3, then interpolate the gaps.
AUM_CATEGORY_BANDS = {
    "Large-Cap":           {"score1_max": 100, "score3_range": (500, 5_000),   "score5_range": (5_000, 40_000)},
    "Mid-Cap":             {"score1_max": 200, "score3_range": (1_000, 10_000), "score5_range": (2_000, 15_000)},
    "Small-Cap":           {"score1_max": 100, "score3_range": (500, 8_000),   "score5_range": (1_000, 10_000)},
    "Flexi-Cap":           {"score1_max": 300, "score3_range": (2_000, 15_000), "score5_range": (5_000, 25_000)},
    "Sectoral / Thematic": {"score1_max": 50,  "score3_range": (200, 3_000),   "score5_range": (500, 5_000)},
    "ELSS":                {"score1_max": 200, "score3_range": (1_000, 10_000), "score5_range": (3_000, 20_000)},
}

# our data's finer-grained "Category" values -> the matrix's 6 AUM buckets.
# Categories the matrix doesn't name (Large & Mid Cap, Multi Cap, Focused,
# Value, Dividend Yield, Contra) fall back to Flexi-Cap's bands as the
# closest generic diversified-equity proxy -- an assumption, not from the file.
AUM_CATEGORY_FALLBACK = "Flexi-Cap"


def resolve_aum_bucket(category):
    """Maps a fund's own category label to one of the matrix's 6 named AUM
    buckets via a case-insensitive substring match (so label variants like
    "Large Cap" / "Large-Cap Fund" / "Large Cap Fund" all resolve the same
    way), falling back to Flexi-Cap for anything genuinely uncovered by the
    matrix (Large & Mid Cap, Multi Cap, Focused, Value, Dividend Yield,
    Contra). Returns (bucket_name, matched) so callers can tell a real match
    from a fallback.

    "Large & Mid Cap" is checked before the plain "large"/"cap" match so it
    doesn't get miscategorized as Large-Cap -- it isn't one of the matrix's
    named buckets and is meant to fall back."""
    if not category:
        return AUM_CATEGORY_FALLBACK, False
    c = category.strip().lower()

    if "sectoral" in c or "thematic" in c:
        return "Sectoral / Thematic", True
    if "elss" in c:
        return "ELSS", True
    if "large" in c and "mid" in c:
        return AUM_CATEGORY_FALLBACK, False
    if "large" in c and "cap" in c:
        return "Large-Cap", True
    if "mid" in c and "cap" in c:
        return "Mid-Cap", True
    if "small" in c and "cap" in c:
        return "Small-Cap", True
    if "flexi" in c:
        return "Flexi-Cap", True
    return AUM_CATEGORY_FALLBACK, False

# (category, column header, source column)
DATASET_COLUMNS = [
    ("Fund Info", "Scheme Code", "__schcode__"),
    ("Fund Info", "Fund Name", "fund_name"),
    ("Fund Info", "Category", "category"),
    ("Fund Info", "Fund Manager", "fund_manager"),
    ("Fund Info", "Benchmark Name", "benchmark_name"),
    ("Return", "Rolling Returns 3Y (%)", "ret_3y"),
    ("Return", "Rolling Returns 5Y (%)", "ret_5y"),
    ("Return", "Alpha (3Y)", "alpha"),
    ("Return", "CAGR vs Category Avg 5Y (%)", "ret_5y"),
    ("Risk", "Standard Deviation (vs Category)", "sd"),
    ("Risk", "Sharpe Ratio (3Y)", "sharpe"),
    ("Risk", "Sortino Ratio (3Y)", "sortino"),
    ("Risk", "Beta", "beta"),
    ("Cost & Operational", "Exit Load Structure", "exit_load_text"),
    ("Cost & Operational", "AUM Size (Cr)", "aum"),
]


# --------------------------------------------------------------------------
# loading + merging
# --------------------------------------------------------------------------

def load_json(path):
    text = Path(path).read_text(encoding="utf-8")
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return json.loads(text, strict=False)


def schcode(value):
    return None if value is None else int(round(float(value)))


def build_intersection_universe(master_payload, risk_payload):
    """master_payload / risk_payload are already-fetched {"data": [...]} dicts --
    the shape load_json() returns for the local snapshots and the shape
    data_sources.py returns for the live API, so this function is indifferent to
    where the rows came from. See load_local_universe() for the file-backed path."""
    master = pd.DataFrame(master_payload["data"])
    master["schcode"] = master["mf_schcode"].apply(schcode)
    master = master.drop_duplicates("schcode").set_index("schcode")

    risk = pd.DataFrame(risk_payload["data"])
    risk["schcode"] = risk["MF_SCHCODE"].apply(schcode)
    risk = risk.drop_duplicates("schcode").set_index("schcode")

    df = master.join(risk, how="left", rsuffix="_risk")
    df = df[df["MainCategory"] == "Equity"].copy()

    df["fund_name"] = df["sch_name"]
    df["category"] = df["Category"]
    df["fund_manager"] = df["FundManager"]
    df["benchmark_name"] = df["BenchmarkName"]
    df["ret_3y"] = df["3YEAR"]
    df["ret_5y"] = df["5YEAR"]
    df["aum"] = df["SchemeAUM"]
    df["exit_load_text"] = df["ExitLoad"]
    for col in ["ALPHA", "BETA", "SD", "SHARPE", "Sortino"]:
        df[col.lower() if col != "Sortino" else "sortino"] = df[col]

    required = ["ret_3y", "ret_5y", "alpha", "sd", "sharpe", "sortino", "beta",
                "exit_load_text", "aum"]
    mask = pd.Series(True, index=df.index)
    for col in required:
        mask &= df[col].notna()

    return df[mask].copy()


def load_local_universe():
    """The original file-backed behaviour, kept for the CLI: read the two static
    snapshots in data/ and build the universe from them."""
    return build_intersection_universe(
        load_json(DATA_DIR / "scheme_masters_1.json"),
        load_json(DATA_DIR / "risk.json"),
    )


def build_dataset_records(df):
    records = []
    for schcode_, row in df.iterrows():
        record = {}
        for _, header, src in DATASET_COLUMNS:
            record[header] = schcode_ if src == "__schcode__" else row.get(src)
        records.append(record)
    records.sort(key=lambda r: (r["AUM Size (Cr)"] is None, -(r["AUM Size (Cr)"] or 0)))
    return records


# --------------------------------------------------------------------------
# scoring helpers -- each mirrors the exact 1/2/3/4/5 bands defined in
# Tanisha Mam's matrix (its "Sheet1"), not a generic z-score or percentile.
# --------------------------------------------------------------------------

def step_score(value, bands):
    """bands: [(min_inclusive, score), ...] sorted descending by min_inclusive,
    with the last entry's min_inclusive = -inf. Returns the score of the first
    band the value clears -- i.e. a direct lookup into the matrix's discrete
    1-5 scoring guide, not an interpolation."""
    if value is None or pd.isna(value):
        return None
    for min_val, score in bands:
        if value >= min_val:
            return float(score)
    return float(bands[-1][1])


def interp_score(value, xp, fp):
    """Continuous stand-in for parameters where the matrix only states 3
    reference points (1/3/5) rather than a full 1-5 breakdown -- e.g. Exit
    Load Structure. Linearly interpolates between those stated points."""
    if value is None or pd.isna(value):
        return None
    return float(np.clip(np.interp(value, xp, fp), 1, 5))


# Alpha (3Y) -- Sheet1: >2%->5, 1.0-2.0%->4, 0-1.0%->3, -1.0-0%->2, <-1.0%->1
ALPHA_BANDS = [(2.0, 5), (1.0, 4), (0.0, 3), (-1.0, 2), (float("-inf"), 1)]

# Sharpe Ratio (3Y) -- Sheet1: >1.20->5, 1.00-1.19->4, 0.70-0.99->3, 0.40-0.69->2, <0.40->1
SHARPE_BANDS = [(1.20, 5), (1.00, 4), (0.70, 3), (0.40, 2), (float("-inf"), 1)]

# Sortino Ratio (3Y) -- Sheet1: >1.50->5, 1.00-1.49->4, 0.70-0.99->3, 0.40-0.69->2, <0.40->1
SORTINO_BANDS = [(1.50, 5), (1.00, 4), (0.70, 3), (0.40, 2), (float("-inf"), 1)]

# Rolling Returns and CAGR each phrase their floor band as "X% or worse" ->
# score 1, i.e. the boundary itself belongs to the bottom tier -- unlike
# Alpha/Sharpe/Sortino, whose floor bands are phrased as a strict "< X".
# step_score()'s >= comparison alone would put exactly -1.0 (or -2.0) into
# the "2" tier instead of "1", so nudge just that tier's threshold up by an
# epsilon to exclude the boundary value from it.
_BOUNDARY_EPS = 1e-9

# Rolling Returns (3Y/5Y) vs Benchmark -- Sheet1's excess-return-vs-benchmark
# bands, applied to (fund's blended 3Y/5Y return) minus (category peer
# average blended return), since no benchmark-index return series exists in
# the source data -- category average is used as the benchmark proxy.
ROLLING_RETURNS_BANDS = [(2.5, 5), (1.5, 4), (0.5, 3), (-1.0 + _BOUNDARY_EPS, 2), (float("-inf"), 1)]

# CAGR vs Category Average (5Y half only -- no 10Y data anywhere) -- Sheet1's
# excess-return-vs-category-average bands, applied to (fund's 5Y return)
# minus (category peer average 5Y return).
CAGR_BANDS = [(3.0, 5), (1.5, 4), (0.5, 3), (-2.0 + _BOUNDARY_EPS, 2), (float("-inf"), 1)]


def score_beta(beta):
    """Beta -- Sheet1: 0.70-0.90->5, 0.91-1.05->4, 1.06-1.20->3, 1.21-1.40->2,
    >1.40->1. Below 0.90 is left open-ended at 5 (even lower market
    sensitivity is still "excellent" per the matrix's own rationale)."""
    if beta is None or pd.isna(beta):
        return None
    if beta <= 0.90:
        return 5.0
    if beta <= 1.05:
        return 4.0
    if beta <= 1.20:
        return 3.0
    if beta <= 1.40:
        return 2.0
    return 1.0


def score_sd_vs_category(sd, cat_mean):
    """Standard Deviation vs Category -- Sheet1 bands are stated as a % of
    category average: >20% below->5, 10-20% below->4, within +-10%->3,
    10-25% above->2, >25% above->1."""
    if sd is None or pd.isna(sd) or cat_mean is None or pd.isna(cat_mean) or cat_mean == 0:
        return None
    pct_diff = (sd - cat_mean) / cat_mean * 100
    if pct_diff <= -20:
        return 5.0
    if pct_diff <= -10:
        return 4.0
    if pct_diff <= 10:
        return 3.0
    if pct_diff <= 25:
        return 2.0
    return 1.0


def parse_exit_load_pct(text):
    if not isinstance(text, str):
        return None
    if re.search(r"\bnil\b", text, re.I):
        return 0.0
    m = re.search(r"(\d+(\.\d+)?)\s*%", text)
    return float(m.group(1)) if m else None


def score_exit_load(text):
    # matrix only gives 3 reference points for this one (1/3/5) -> interpolate
    pct = parse_exit_load_pct(text)
    if pct is None:
        return None
    return interp_score(pct, [0, 1, 2], [5, 3, 1])


def score_aum(aum, bands):
    """AUM Size -- category-specific Cr ranges straight from the matrix's
    'AUM Size Scoring Guide by Fund Category' table (see AUM_CATEGORY_BANDS).
    Inside the score-5 range scores 5; inside score-3 (and not score-5) scores
    3; below the score-1 ceiling scores 1; gaps between checkpoints are
    interpolated. Above the score-5 range there's no explicit matrix band, so
    it's tapered back down over a span equal to the score-5 range's own width
    -- consistent with the matrix's general "too large for category" = 1
    principle, but this specific falloff shape is our extrapolation, not hers."""
    if aum is None or pd.isna(aum):
        return None
    lo1 = bands["score1_max"]
    lo3, hi3 = bands["score3_range"]
    lo5, hi5 = bands["score5_range"]

    if lo5 <= aum <= hi5:
        return 5.0
    if aum > hi5:
        span = max(hi5 - lo5, 1)
        return float(np.clip(5 - 4 * ((aum - hi5) / span), 1, 5))
    if aum < lo1:
        return 1.0
    if lo3 <= aum <= hi3:
        return 3.0
    if aum < lo3:
        return float(np.interp(aum, [lo1, lo3], [1, 3]))
    return float(np.interp(aum, [hi3, lo5], [3, 5]))  # gap between score-3 and score-5 ranges


def compute_category_stats(df):
    stats = {}
    df = df.copy()
    df["ret_blend"] = df[["ret_3y", "ret_5y"]].mean(axis=1)
    for cat, g in df.groupby("category"):
        stats[cat] = {
            "ret_blend_mean": g["ret_blend"].mean(),
            "ret_5y_mean": g["ret_5y"].mean(),
            "sd_mean": g["sd"].mean(),
        }
    return stats


def score_fund(row, cat_stats):
    cat = cat_stats.get(row["category"], {})
    scores = {}

    blended_return = np.nanmean([row["ret_3y"], row["ret_5y"]])
    cat_ret_blend_mean = cat.get("ret_blend_mean")
    excess_return = None if cat_ret_blend_mean is None else blended_return - cat_ret_blend_mean
    scores["rolling_returns_vs_benchmark"] = step_score(excess_return, ROLLING_RETURNS_BANDS)

    scores["alpha_3y"] = step_score(row["alpha"], ALPHA_BANDS)

    cat_ret_5y_mean = cat.get("ret_5y_mean")
    excess_5y = None if cat_ret_5y_mean is None else row["ret_5y"] - cat_ret_5y_mean
    scores["cagr_vs_category_avg"] = step_score(excess_5y, CAGR_BANDS)

    scores["sd_vs_category"] = score_sd_vs_category(row["sd"], cat.get("sd_mean"))
    scores["sharpe_3y"] = step_score(row["sharpe"], SHARPE_BANDS)
    scores["sortino_3y"] = step_score(row["sortino"], SORTINO_BANDS)
    scores["beta"] = score_beta(row["beta"])

    scores["exit_load_structure"] = score_exit_load(row["exit_load_text"])
    aum_bucket, _ = resolve_aum_bucket(row["category"])
    scores["aum_size_category_adjusted"] = score_aum(row["aum"], AUM_CATEGORY_BANDS[aum_bucket])

    return scores


def composite_and_rating(param_scores):
    available = {k: v for k, v in param_scores.items() if v is not None}
    total_weight = sum(WEIGHTS.values())
    available_weight = sum(WEIGHTS[k] for k in available)
    if available_weight == 0:
        return None, 0.0, "Insufficient Data", "N/A"

    composite = sum(WEIGHTS[k] * v for k, v in available.items()) / available_weight
    completeness = available_weight / total_weight

    for cutoff, rating, rec in RATING_BANDS:
        if composite >= cutoff:
            return round(composite, 2), round(completeness, 2), rating, rec
    return round(composite, 2), round(completeness, 2), "Poor", "Reject"


def category_rollup(param_scores):
    rollup = {}
    for cat in set(PARAM_CATEGORY.values()):
        keys = [k for k, c in PARAM_CATEGORY.items() if c == cat]
        vals = [param_scores[k] for k in keys if param_scores[k] is not None]
        rollup[cat] = round(float(np.mean(vals)), 2) if vals else None
    return rollup


def build_score_rows(df, top_n=None):
    cat_stats = compute_category_stats(df)
    rows = []
    for schcode_, row in df.iterrows():
        param_scores = score_fund(row, cat_stats)
        composite, completeness, rating, rec = composite_and_rating(param_scores)
        if composite is None:
            continue
        rows.append({
            "schcode": schcode_,
            "fund": row["fund_name"],
            "category": row["category"],
            "composite": composite,
            "data_completeness": completeness,
            "rating": rating,
            "recommendation": rec,
            "category_scores": category_rollup(param_scores),
            "parameter_scores": {k: (round(v, 2) if v is not None else None) for k, v in param_scores.items()},
        })

    rows.sort(key=lambda r: r["composite"], reverse=True)
    for i, r in enumerate(rows, start=1):
        r["rank"] = i

    if top_n:
        rows = rows[:top_n]
    return rows


# --------------------------------------------------------------------------
# output writers
# --------------------------------------------------------------------------

def clean(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, float):
        return round(value, 4)
    return value


def write_dataset_json(records, path):
    cleaned = [{k: clean(v) for k, v in r.items()} for r in records]
    path.write_text(json.dumps({"count": len(cleaned), "data": cleaned}, indent=2, default=str), encoding="utf-8")


def write_grouped_xlsx(path, sheet_title, columns, rows_getter, n_rows):
    """columns: list of (category, header); rows_getter(row_index) -> list of cell values in column order"""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    cat_fill = PatternFill("solid", fgColor="2E75B6")
    cat_font = Font(color="FFFFFF", bold=True)

    start = 1
    prev_cat = columns[0][0]
    for i, (cat, _) in enumerate(columns + [(None, None)], start=1):
        if cat != prev_cat:
            end = i - 1
            ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
            cell = ws.cell(row=1, column=start, value=prev_cat.upper())
            cell.fill = cat_fill
            cell.font = cat_font
            cell.alignment = Alignment(horizontal="center")
            start = i
            prev_cat = cat

    for i, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=2, column=i, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    for r in range(n_rows):
        values = rows_getter(r)
        for c, value in enumerate(values, start=1):
            ws.cell(row=r + 3, column=c, value=value)

    for i, (_, header) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, min(32, len(header) // 1.3 + 4))

    ws.freeze_panes = "C3"
    wb.save(path)


def write_dataset_xlsx(records, path):
    columns = [(cat, header) for cat, header, _ in DATASET_COLUMNS]
    write_grouped_xlsx(
        path, "Intersection Fund Dataset", columns,
        lambda r: [clean(records[r][header]) for _, header, _ in DATASET_COLUMNS],
        len(records),
    )


def write_scores_json(rows, path):
    path.write_text(json.dumps({"count": len(rows), "data": rows}, indent=2, default=str), encoding="utf-8")


def write_scores_xlsx(rows, path):
    id_cols = ["rank", "schcode", "fund", "category", "composite", "data_completeness", "rating", "recommendation"]
    id_headers = ["Rank", "Scheme Code", "Fund Name", "Category", "Composite Score",
                  "Data Completeness", "Rating", "Recommendation"]
    param_keys = list(WEIGHTS.keys())
    columns = [("Fund Info", h) for h in id_headers] + [(PARAM_CATEGORY[k], PARAM_LABELS[k]) for k in param_keys]

    def row_values(r):
        row = rows[r]
        values = [row[k] for k in id_cols]
        values += [row["parameter_scores"][k] for k in param_keys]
        return values

    write_grouped_xlsx(path, "Intersection Fund Scores", columns, row_values, len(rows))


def print_aum_category_diagnostic(df):
    """One-time visibility check: which AUM bucket every distinct category in
    the universe actually resolves to, so a category-label mismatch (e.g. a
    future data refresh renaming "Large Cap Fund" to something else) shows up
    here as an unexpected fallback instead of silently scoring AUM wrong."""
    print("AUM category -> bucket mapping:")
    for cat in sorted(df["category"].dropna().unique()):
        bucket, matched = resolve_aum_bucket(cat)
        tag = bucket if matched else f"{bucket} (fallback)"
        print(f"  {cat!r:30s} -> {tag}")


# --------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--top", type=int, default=None, help="only keep the top N funds in the score output")
    parser.add_argument("--out-suffix", default="",
                         help="append to output basenames, e.g. _v2 -> equity_fund_intersection_dataset_v2.xlsx "
                              "(handy when the default files are open/locked in Excel)")
    args = parser.parse_args()

    df = load_local_universe()
    print(f"Intersection universe: {len(df)} Equity funds have all 9 parameters present")
    print_aum_category_diagnostic(df)

    records = build_dataset_records(df)
    write_dataset_json(records, OUTPUT_DIR / f"equity_fund_intersection_dataset{args.out_suffix}.json")
    write_dataset_xlsx(records, OUTPUT_DIR / f"equity_fund_intersection_dataset{args.out_suffix}.xlsx")
    print(f"Wrote {len(records)} funds -> equity_fund_intersection_dataset{args.out_suffix}.json / .xlsx")

    rows = build_score_rows(df, top_n=args.top)
    write_scores_json(rows, OUTPUT_DIR / f"equity_fund_intersection_scores{args.out_suffix}.json")
    write_scores_xlsx(rows, OUTPUT_DIR / f"equity_fund_intersection_scores{args.out_suffix}.xlsx")
    print(f"Scored {len(rows)} funds -> equity_fund_intersection_scores{args.out_suffix}.json / .xlsx")


if __name__ == "__main__":
    main()
